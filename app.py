# -*- coding: utf-8 -*-
"""
Report Semanal - servidor local
--------------------------------
Abre a interface de edição no navegador e, ao gerar a apresentação:
  1. guarda os dados da semana em  dados/semanas/AAAA-MM-DD.json
  2. grava a apresentação em        saida/AAAA/AAAA-MM-DD/report_AAAA-MM-DD.html
  3. gera o PDF com o Edge (ou Chrome) em modo headless
  4. reconstrói o historico.xlsx a partir de todas as semanas salvas

Só usa a biblioteca padrão do Python. O pacote openpyxl é opcional:
sem ele, o histórico é gravado em CSV (pasta historico_csv).
O servidor escuta apenas em 127.0.0.1 (só esta máquina acessa).
"""
import csv
import getpass
import uuid
import datetime as dt
import json
import os
import re
import shutil
import socket
import subprocess
import sys
import tempfile
import threading
import time
import webbrowser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    TEM_XLSX = True
except ImportError:
    TEM_XLSX = False

VERSAO = "1.1"
BASE = Path(__file__).resolve().parent
WEB = BASE / "web"
DADOS = BASE / "dados"
SEMANAS = DADOS / "semanas"
SAIDA = BASE / "saida"
RASCUNHO = DADOS / "rascunho.json"
EXEMPLO = DADOS / "exemplo.json"
HIST_XLSX = BASE / "historico.xlsx"
HIST_CSV = BASE / "historico_csv"
LOCK = DADOS / "em_uso.json"           # quem está editando agora (uso em pasta compartilhada)
CONFIG_LOCAL = BASE / "config_local.json"   # opcional: {"navegador": "C:\\...\\msedge.exe", "porta": 8765}

USUARIO = os.environ.get("USERNAME") or os.environ.get("USER") or ""
try:
    USUARIO = USUARIO or getpass.getuser()
except Exception:
    USUARIO = USUARIO or "desconhecido"
MAQUINA = socket.gethostname()
SESSAO = uuid.uuid4().hex
LOCK_VALIDADE = 120      # segundos sem sinal para considerar a trava abandonada
LOCK_BATIDA = 30         # de quanto em quanto tempo esta sessão renova a trava

RE_DATA = re.compile(r"^\d{4}-\d{2}-\d{2}$")
TIPOS = {".html": "text/html; charset=utf-8", ".js": "text/javascript; charset=utf-8",
         ".css": "text/css; charset=utf-8", ".json": "application/json; charset=utf-8",
         ".png": "image/png", ".svg": "image/svg+xml", ".ico": "image/x-icon"}


# --------------------------------------------------------------------------- utilitários
def ler_json(p, padrao=None):
    try:
        return json.loads(Path(p).read_text(encoding="utf-8"))
    except Exception:
        return padrao


def gravar_json(p, obj):
    """Grava de forma atômica (arquivo temporário + rename) para não corromper em caso de erro."""
    p = Path(p)
    p.parent.mkdir(parents=True, exist_ok=True)
    tmp = p.with_suffix(p.suffix + ".tmp")
    tmp.write_text(json.dumps(obj, ensure_ascii=False, indent=1), encoding="utf-8")
    os.replace(tmp, p)


def config_local():
    return ler_json(CONFIG_LOCAL, {}) or {}


def para_data(s):
    try:
        return dt.date.fromisoformat(str(s)[:10]) if s else None
    except ValueError:
        return None


def lista_semanas():
    if not SEMANAS.exists():
        return []
    out = []
    for f in sorted(SEMANAS.glob("*.json"), reverse=True):
        d = f.stem
        if not RE_DATA.match(d):
            continue
        pasta = SAIDA / d[:4] / d
        out.append({"data": d,
                    "pdf": (pasta / f"report_{d}.pdf").exists(),
                    "html": (pasta / f"report_{d}.html").exists()})
    return out


def abrir_no_sistema(caminho):
    caminho = str(caminho)
    if sys.platform.startswith("win"):
        os.startfile(caminho)  # noqa
    elif sys.platform == "darwin":
        subprocess.Popen(["open", caminho])
    else:
        subprocess.Popen(["xdg-open", caminho])


# --------------------------------------------------------------------------- regras (mesmas do render.js)
STATUS_VALIDOS = ["Concluído", "Não iniciado", "No prazo", "Atenção", "Atrasado", "Bloqueado"]


def status_de(a, hoje):
    if a.get("status") in STATUS_VALIDOS:
        return a["status"]
    try:
        pct = float(a.get("pct") or 0)
    except (TypeError, ValueError):
        pct = 0
    if pct >= 100:
        return "Concluído"
    ini, fim = para_data(a.get("inicio")), para_data(a.get("fim"))
    if not ini or not fim:
        return "Não iniciado"
    if fim < hoje:
        return "Atrasado"
    if ini > hoje:
        return "Não iniciado"
    esperado = 100 * (hoje - ini).days / ((fim - ini).days + 1)
    return "Atenção" if pct < esperado - 25 else "No prazo"


def sprint_da_data(dados, hoje):
    for s in dados.get("sprints") or []:
        i, f = para_data(s.get("inicio")), para_data(s.get("fim"))
        if i and f and i <= hoje <= f:
            return f"{str(s.get('trimestre', '')).split(' ')[0]} · Sprint {s.get('numero', '')}"
    return ""


# --------------------------------------------------------------------------- trava de edição
_trava = threading.Lock()


def ler_lock():
    l = ler_json(LOCK)
    return l if isinstance(l, dict) else None


def lock_ativo(l):
    return bool(l) and (time.time() - float(l.get("batida", 0))) < LOCK_VALIDADE


def sou_dono():
    l = ler_lock()
    return bool(l) and l.get("sessao") == SESSAO


def pegar_lock(forcar=False):
    """Tenta assumir a edição. Devolve None se conseguiu, ou os dados de quem está editando."""
    with _trava:
        l = ler_lock()
        if lock_ativo(l) and l.get("sessao") != SESSAO and not forcar:
            return l
        agora = time.time()
        gravar_json(LOCK, {"usuario": USUARIO, "maquina": MAQUINA, "sessao": SESSAO,
                           "desde": l.get("desde", agora) if l and l.get("sessao") == SESSAO else agora,
                           "batida": agora})
        return None


def renovar_lock():
    while True:
        time.sleep(LOCK_BATIDA)
        try:
            with _trava:
                l = ler_lock()
                if l and l.get("sessao") == SESSAO:
                    l["batida"] = time.time()
                    gravar_json(LOCK, l)
        except Exception:
            pass


def soltar_lock():
    try:
        with _trava:
            if sou_dono():
                LOCK.unlink()
    except Exception:
        pass


def info_lock():
    l = ler_lock()
    if not lock_ativo(l) or l.get("sessao") == SESSAO:
        return None
    return {"usuario": l.get("usuario"), "maquina": l.get("maquina"),
            "desde": dt.datetime.fromtimestamp(float(l.get("desde", 0))).strftime("%d/%m %H:%M")}


def carimbo(dados):
    dados["_meta"] = {"usuario": USUARIO, "maquina": MAQUINA,
                      "quando": dt.datetime.now().strftime("%Y-%m-%d %H:%M")}
    return dados


# --------------------------------------------------------------------------- PDF
def achar_navegador():
    cfg = config_local().get("navegador") or os.environ.get("REPORT_NAVEGADOR")
    if cfg and Path(cfg).exists():
        return cfg
    candidatos = []
    for var in ("PROGRAMFILES(X86)", "PROGRAMFILES", "LOCALAPPDATA"):
        raiz = os.environ.get(var)
        if raiz:
            candidatos += [Path(raiz) / "Microsoft/Edge/Application/msedge.exe",
                           Path(raiz) / "Google/Chrome/Application/chrome.exe"]
    for c in candidatos:
        if c.exists():
            return str(c)
    for nome in ("msedge", "microsoft-edge", "google-chrome", "chrome", "chromium", "chromium-browser"):
        w = shutil.which(nome)
        if w:
            return w
    return None


def gerar_pdf(html_path, pdf_path):
    """Imprime o HTML em PDF com o navegador em modo headless. Devolve None se deu certo, ou a mensagem de erro."""
    nav = achar_navegador()
    if not nav:
        return ("Não encontrei o Edge nem o Chrome. Abra o HTML gerado e use Ctrl+P > Salvar como PDF, "
                "ou informe o caminho do navegador em config_local.json.")
    pdf_path = Path(pdf_path)
    if pdf_path.exists():
        try:
            pdf_path.unlink()
        except PermissionError:
            return "O PDF desta semana está aberto em outro programa. Feche-o e gere novamente."
    ultimo_erro = ""
    for modo in ("--headless=new", "--headless"):
        perfil = tempfile.mkdtemp(prefix="report_pdf_")
        cmd = [nav, modo, "--disable-gpu", "--no-first-run", "--no-default-browser-check",
               "--disable-extensions", f"--user-data-dir={perfil}",
               "--no-pdf-header-footer", "--print-to-pdf-no-header",
               "--run-all-compositor-stages-before-draw", "--virtual-time-budget=4000",
               f"--print-to-pdf={pdf_path}", Path(html_path).resolve().as_uri()]
        if not sys.platform.startswith("win"):
            cmd.insert(2, "--no-sandbox")
        try:
            r = subprocess.run(cmd, capture_output=True, timeout=60,
                               creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
            ultimo_erro = (r.stderr or b"").decode("utf-8", "replace")[-400:]
        except subprocess.TimeoutExpired:
            ultimo_erro = "o navegador demorou demais para responder"
        except OSError as e:
            ultimo_erro = str(e)
        finally:
            time.sleep(0.3)
            shutil.rmtree(perfil, ignore_errors=True)
        if pdf_path.exists() and pdf_path.stat().st_size > 1000:
            return None
    return f"O PDF não foi gerado ({ultimo_erro.strip() or 'sem detalhes'}). O HTML foi gravado normalmente."


# --------------------------------------------------------------------------- histórico (Excel ou CSV)
def linhas_historico():
    """Monta as tabelas do histórico a partir de todas as semanas salvas (fonte da verdade: os JSON)."""
    tab = {
        "Atividades": [["Semana", "Sprint", "Frente", "Atividade", "Início", "Fim", "% concluído",
                        "% semana anterior", "Status", "Status manual?", "Marco?"]],
        "Bloqueios": [["Semana", "Sprint", "Frente", "Bloqueio / dependência", "Área", "Desde",
                       "Dias em aberto", "Impacto", "Pedido / ação", "Criticidade"]],
        "Semana": [["Semana", "Sprint", "Destaques", "Próximos passos", "Gerado por", "Gerado em"]],
        "Comunicações": [["Semana", "Ordem", "Incluída?", "Título", "Subtítulo", "Posição", "Texto", "Tem imagem?"]],
        "Frentes": [["Sigla", "Nome", "Responsável", "Objetivo"]],
        "Sprints": [["Trimestre", "Nº", "Início", "Fim"]],
    }
    semanas = sorted(SEMANAS.glob("*.json")) if SEMANAS.exists() else []
    ultimo = None
    for f in semanas:
        if not RE_DATA.match(f.stem):
            continue
        dados = ler_json(f)
        if not dados:
            continue
        ultimo = dados
        hoje = dt.date.fromisoformat(f.stem)
        sp = sprint_da_data(dados, hoje)
        nomes = {x.get("id"): x.get("nome") for x in dados.get("frentes") or []}
        for a in dados.get("atividades") or []:
            if not a.get("nome"):
                continue
            marco = bool(a.get("marco"))
            tab["Atividades"].append([
                hoje, sp, nomes.get(a.get("frente"), a.get("frente")), a.get("nome"),
                para_data(a.get("inicio")), None if marco else para_data(a.get("fim")),
                None if marco else _num(a.get("pct")), None if marco else _num(a.get("pctAnterior")),
                "Marco" if marco else status_de(a, hoje),
                "" if marco else ("Sim" if a.get("status") in STATUS_VALIDOS else "Não"),
                "Sim" if marco else "Não"])
        for b in dados.get("bloqueios") or []:
            if not b.get("descricao"):
                continue
            desde = para_data(b.get("desde"))
            tab["Bloqueios"].append([hoje, sp, nomes.get(b.get("frente"), b.get("frente")), b.get("descricao"),
                                     b.get("area"), desde, (hoje - desde).days if desde else None,
                                     b.get("impacto"), b.get("pedido"), b.get("criticidade")])
        sem = dados.get("semana") or {}
        meta = dados.get("_meta") or {}
        tab["Semana"].append([hoje, sp, sem.get("destaques", ""), sem.get("proximos", ""),
                              meta.get("usuario", ""), meta.get("quando", "")])
        for i, c in enumerate(dados.get("comunicacoes") or [], 1):
            tab["Comunicações"].append([hoje, i, "Sim" if c.get("incluir") else "Não", c.get("titulo"),
                                        c.get("subtitulo"), c.get("posicao"), c.get("texto"),
                                        "Sim" if c.get("imagem") else "Não"])
    if ultimo:
        for x in ultimo.get("frentes") or []:
            tab["Frentes"].append([x.get("id"), x.get("nome"), x.get("responsavel"), x.get("objetivo")])
        for s in ultimo.get("sprints") or []:
            tab["Sprints"].append([s.get("trimestre"), s.get("numero"), para_data(s.get("inicio")), para_data(s.get("fim"))])
    return tab


def _num(v):
    try:
        return None if v in ("", None) else float(v)
    except (TypeError, ValueError):
        return None


def gravar_historico():
    """Reconstrói o histórico inteiro. Devolve (caminho, aviso_ou_None)."""
    tab = linhas_historico()
    if TEM_XLSX:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        cab_fill = PatternFill("solid", fgColor="EEEEEE")
        for nome, linhas in tab.items():
            ws = wb.create_sheet(nome)
            for ln in linhas:
                ws.append(ln)
            for c in ws[1]:
                c.font = Font(bold=True)
                c.fill = cab_fill
            ws.freeze_panes = "A2"
            if len(linhas) > 1:
                ws.auto_filter.ref = ws.dimensions
            for col in range(1, len(linhas[0]) + 1):
                larg = max(len(str(l[col - 1])) if l[col - 1] is not None else 0 for l in linhas[:200])
                ws.column_dimensions[get_column_letter(col)].width = min(60, max(10, larg + 2))
                for cel in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                    v = cel[0].value
                    if isinstance(v, dt.date):
                        cel[0].number_format = "DD/MM/YYYY"
                    elif isinstance(v, str) and "\n" in v:
                        cel[0].alignment = Alignment(wrap_text=True, vertical="top")
        try:
            tmp = HIST_XLSX.with_name("~historico_tmp.xlsx")
            wb.save(tmp)
            os.replace(tmp, HIST_XLSX)
            return str(HIST_XLSX), None
        except PermissionError:
            try:
                tmp.unlink()
            except Exception:
                pass
            return str(HIST_XLSX), ("O historico.xlsx está aberto no Excel e não pôde ser atualizado. "
                                    "Feche-o e clique em Gerar de novo (os dados da semana já estão salvos).")
    HIST_CSV.mkdir(exist_ok=True)
    aviso = None
    for nome, linhas in tab.items():
        arq = HIST_CSV / (nome.lower().replace("ç", "c").replace("õ", "o") + ".csv")
        try:
            with open(arq, "w", newline="", encoding="utf-8-sig") as fh:
                w = csv.writer(fh, delimiter=";")
                for ln in linhas:
                    w.writerow([v.strftime("%d/%m/%Y") if isinstance(v, dt.date) else
                                (str(v).replace(".", ",") if isinstance(v, float) else v) for v in ln])
        except PermissionError:
            aviso = f"{arq.name} está aberto em outro programa e não foi atualizado."
    return str(HIST_CSV), aviso


# --------------------------------------------------------------------------- ações da API
def estado():
    dados, origem = ler_json(RASCUNHO), "rascunho"
    if not dados:
        sem = lista_semanas()
        if sem:
            dados, origem = ler_json(SEMANAS / f"{sem[0]['data']}.json"), "semana " + sem[0]["data"]
    if not dados:
        dados, origem = ler_json(EXEMPLO, {}), "exemplo"
    if not sou_dono():
        pegar_lock()          # se a trava estiver livre (ou abandonada), esta sessão assume
    return {"dados": dados, "origem": origem, "semanas": lista_semanas(), "xlsx": TEM_XLSX,
            "navegador": achar_navegador(), "pasta": str(BASE), "versao": VERSAO,
            "usuario": USUARIO, "modo": "edicao" if sou_dono() else "leitura", "lock": info_lock(),
            "meta": (dados or {}).get("_meta")}


def gerar(payload):
    dados, html = payload.get("dados"), payload.get("html")
    if not isinstance(dados, dict) or not isinstance(html, str) or not html.startswith("<!DOCTYPE html>"):
        raise ValueError("Dados inválidos recebidos da interface.")
    data = str((dados.get("config") or {}).get("dataReferencia", ""))
    if not RE_DATA.match(data):
        raise ValueError("Data de referência inválida.")
    avisos = []
    carimbo(dados)
    gravar_json(SEMANAS / f"{data}.json", dados)
    gravar_json(RASCUNHO, dados)
    pasta = SAIDA / data[:4] / data
    pasta.mkdir(parents=True, exist_ok=True)
    html_path = pasta / f"report_{data}.html"
    pdf_path = pasta / f"report_{data}.pdf"
    html_path.write_text(html, encoding="utf-8")
    erro_pdf = gerar_pdf(html_path, pdf_path)
    if erro_pdf:
        avisos.append(erro_pdf)
    hist, aviso_hist = gravar_historico()
    if aviso_hist:
        avisos.append(aviso_hist)
    return {"ok": True, "data": data, "pasta": str(pasta), "html": str(html_path),
            "pdf": None if erro_pdf else str(pdf_path), "historico": hist, "avisos": avisos,
            "semanas": lista_semanas(), "meta": dados["_meta"]}


def caminho_para_abrir(alvo, data):
    if alvo == "raiz":
        return BASE
    if alvo == "historico":
        return HIST_XLSX if TEM_XLSX else HIST_CSV
    if not RE_DATA.match(data or ""):
        raise ValueError("Data inválida.")
    pasta = SAIDA / data[:4] / data
    return {"pasta": pasta, "pdf": pasta / f"report_{data}.pdf", "html": pasta / f"report_{data}.html"}[alvo]


# --------------------------------------------------------------------------- servidor HTTP
class Handler(BaseHTTPRequestHandler):
    server_version = "ReportSemanal/" + VERSAO

    def log_message(self, fmt, *args):  # silencioso, exceto erros
        pass

    def _resposta(self, codigo, corpo, tipo="application/json; charset=utf-8"):
        if isinstance(corpo, (dict, list)):
            corpo = json.dumps(corpo, ensure_ascii=False, default=str).encode("utf-8")
        elif isinstance(corpo, str):
            corpo = corpo.encode("utf-8")
        self.send_response(codigo)
        self.send_header("Content-Type", tipo)
        self.send_header("Content-Length", str(len(corpo)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(corpo)

    def _corpo(self):
        n = int(self.headers.get("Content-Length") or 0)
        if n > 200 * 1024 * 1024:
            raise ValueError("Conteúdo grande demais (verifique o tamanho das imagens).")
        return json.loads(self.rfile.read(n).decode("utf-8")) if n else {}

    def _origem_ok(self):
        # só aceita chamadas vindas da própria interface (proteção contra outros sites)
        o = self.headers.get("Origin")
        return o is None or o == f"http://127.0.0.1:{self.server.server_port}" or o == f"http://localhost:{self.server.server_port}"

    def do_GET(self):
        u = urlparse(self.path)
        try:
            if u.path == "/api/estado":
                return self._resposta(200, estado())
            if u.path == "/api/semana":
                d = (parse_qs(u.query).get("data") or [""])[0]
                if not RE_DATA.match(d) or not (SEMANAS / f"{d}.json").exists():
                    return self._resposta(404, {"erro": "Semana não encontrada."})
                return self._resposta(200, {"dados": ler_json(SEMANAS / f"{d}.json")})
            rel = "index.html" if u.path in ("/", "") else u.path.lstrip("/")
            arq = (WEB / rel).resolve()
            if WEB.resolve() not in arq.parents or not arq.is_file():
                return self._resposta(404, "Não encontrado", "text/plain; charset=utf-8")
            return self._resposta(200, arq.read_bytes(), TIPOS.get(arq.suffix.lower(), "application/octet-stream"))
        except Exception as e:  # noqa
            return self._resposta(500, {"erro": str(e)})

    def do_POST(self):
        if not self._origem_ok():
            return self._resposta(403, {"erro": "Origem não permitida."})
        u = urlparse(self.path)
        try:
            corpo = self._corpo()
            if u.path in ("/api/salvar", "/api/gerar") and not sou_dono():
                quem = info_lock() or {}
                return self._resposta(423, {"erro": f"{quem.get('usuario', 'Outra pessoa')} está editando a apresentação agora. "
                                                    "Suas alterações não foram salvas.", "lock": quem})
            if u.path == "/api/assumir":
                pegar_lock(forcar=bool(corpo.get("forcar")))
                e = estado()
                return self._resposta(200, e)
            if u.path == "/api/salvar":
                if not isinstance(corpo.get("dados"), dict):
                    raise ValueError("Dados inválidos.")
                gravar_json(RASCUNHO, carimbo(corpo["dados"]))
                return self._resposta(200, {"ok": True, "hora": dt.datetime.now().strftime("%H:%M")})
            if u.path == "/api/gerar":
                print(f"  Gerando apresentação de {corpo.get('dados', {}).get('config', {}).get('dataReferencia')}...")
                r = gerar(corpo)
                print(f"  OK: {r['pasta']}" + (f"  (avisos: {len(r['avisos'])})" if r["avisos"] else ""))
                return self._resposta(200, r)
            if u.path == "/api/abrir":
                alvo = caminho_para_abrir(corpo.get("alvo"), corpo.get("data"))
                if not Path(alvo).exists():
                    return self._resposta(404, {"erro": "Arquivo ainda não existe."})
                abrir_no_sistema(alvo)
                return self._resposta(200, {"ok": True})
            if u.path == "/api/encerrar":
                soltar_lock()
                self._resposta(200, {"ok": True})
                threading.Thread(target=self.server.shutdown, daemon=True).start()
                return
            return self._resposta(404, {"erro": "Ação desconhecida."})
        except ValueError as e:
            return self._resposta(400, {"erro": str(e)})
        except Exception as e:  # noqa
            print("  ERRO:", repr(e))
            return self._resposta(500, {"erro": f"Erro inesperado: {e}"})


def porta_livre(preferida):
    for p in range(preferida, preferida + 30):
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            try:
                s.bind(("127.0.0.1", p))
                return p
            except OSError:
                continue
    raise SystemExit("Nenhuma porta livre encontrada.")


def main():
    for p in (DADOS, SEMANAS, SAIDA):
        p.mkdir(parents=True, exist_ok=True)
    porta = porta_livre(int(config_local().get("porta", 8765)))
    srv = ThreadingHTTPServer(("127.0.0.1", porta), Handler)
    url = f"http://127.0.0.1:{porta}/"
    nav = achar_navegador()
    print("=" * 64)
    print(f"  Report Semanal v{VERSAO}")
    print(f"  Interface: {url}")
    print(f"  Pasta:     {BASE}")
    print(f"  Histórico: {'Excel (historico.xlsx)' if TEM_XLSX else 'CSV (openpyxl não instalado)'}")
    print(f"  PDF:       {nav or 'navegador não encontrado - PDF desativado'}")
    print(f"  Usuário:   {USUARIO} @ {MAQUINA}")
    print("  Para encerrar: botão 'Encerrar' na interface, ou feche esta janela.")
    print("=" * 64)
    quem = pegar_lock()
    if quem:
        print(f"  ATENÇÃO: {quem.get('usuario')} ({quem.get('maquina')}) está editando agora. Abrindo em modo leitura.")
    threading.Thread(target=renovar_lock, daemon=True).start()
    if "--sem-navegador" not in sys.argv:
        threading.Timer(0.6, lambda: webbrowser.open(url)).start()
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        soltar_lock()
        srv.server_close()
        print("  Encerrado.")


if __name__ == "__main__":
    main()
